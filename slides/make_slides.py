#!/usr/bin/env python3
"""Generate the first half (through the algorithms) of the FMCAD 2026 talk
"Parallelizing Congruence Closure".  Run:  python3 make_slides.py
Output: parallelcc-part1.pptx next to this script."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
import os

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x1F, 0x2A, 0x44)
INK    = RGBColor(0x2B, 0x2F, 0x36)
GRAY   = RGBColor(0x6B, 0x71, 0x7A)
LIGHT  = RGBColor(0xC9, 0xCD, 0xD3)
ORANGE = RGBColor(0xC8, 0x50, 0x1E)
AMBER  = RGBColor(0xE0, 0xA3, 0x00)
AMBER_FILL = RGBColor(0xFF, 0xF3, 0xCC)
PURPLE = RGBColor(0x7B, 0x4F, 0xB5)
RED    = RGBColor(0xC0, 0x39, 0x2B)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
EDGE   = RGBColor(0x8A, 0x8F, 0x98)

CLASS_STYLE = {   # fill, outline
    'leaf': (RGBColor(0xEE, 0xF0, 0xF3), RGBColor(0x9A, 0xA0, 0xA8)),
    'a':    (RGBColor(0xDC, 0xEA, 0xF8), RGBColor(0x2E, 0x6F, 0xB0)),
    'x':    (RGBColor(0xDD, 0xF1, 0xE2), RGBColor(0x3A, 0x9A, 0x5B)),
    'm':    (RGBColor(0xFB, 0xE4, 0xD3), ORANGE),
}

FONT = 'Calibri'
MONO = 'Consolas'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_NO = [0]

# ---------------------------------------------------------------- helpers
def no_shadow(shape):
    shape.shadow.inherit = False
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)

def set_arrow(conn, head=True):
    ln = conn.line._get_or_add_ln()
    tag = 'a:tailEnd' if head else 'a:headEnd'
    for el in ln.findall(qn(tag)):
        ln.remove(el)
    ln.append(ln.makeelement(qn(tag), {'type': 'triangle', 'w': 'med', 'len': 'med'}))

def style_run(run, size, bold=False, color=INK, font=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color

def add_runs(p, text, size, color=INK, font=FONT, bold=False):
    """'**' toggles bold inside text."""
    parts = text.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        style_run(r, size, bold=(bold or i % 2 == 1), color=color, font=font)

def add_text(slide, text, x, y, w, h, size=20, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, italic=False,
             line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        add_runs(p, line, size, color=color, font=font, bold=bold)
        if italic:
            for r in p.runs:
                r.font.italic = True
    return tb

def set_bullet(p, level, char='•', color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(0.30 + 0.34 * level))))
    pPr.set('indent', str(-int(Inches(0.26))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buClr'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if color is not None:
        buClr = pPr.makeelement(qn('a:buClr'), {})
        srgb = buClr.makeelement(qn('a:srgbClr'), {'val': '%02X%02X%02X' % (color[0], color[1], color[2])})
        buClr.append(srgb)
        pPr.append(buClr)
    pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char}))

def add_bullets(slide, items, x, y, w, h, size=20, color=INK, gap=6, font=FONT):
    """items: list of str or (level, str).  '**' toggles bold.  A leading
    '#' marks a heading line (bold, no bullet)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    first = True
    for it in items:
        level, text = (it if isinstance(it, tuple) else (0, it))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if text.startswith('#'):
            p.space_before = Pt(gap + 6)
            p.space_after = Pt(2)
            add_runs(p, text[1:].strip(), size, color=NAVY, font=font, bold=True)
            continue
        p.space_before = Pt(gap if level == 0 else 2)
        lvl_size = size if level == 0 else size - 3
        set_bullet(p, level, char='•' if level == 0 else '–', color=GRAY)
        add_runs(p, text, lvl_size, color=color, font=font)
    return tb

def add_rect(slide, x, y, w, h, fill=None, line=None, width=1.0, dash=None,
             rounded=True, radius=0.15, text=None, size=14, bold=False,
             color=INK, align=PP_ALIGN.CENTER, font=FONT):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    no_shadow(s)
    if rounded:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
        if dash:
            s.line.dash_style = dash
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if text is not None:
        lines = text.split('\n')
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            add_runs(p, line, size, color=color, font=font, bold=bold)
    return s

def add_line(slide, x1, y1, x2, y2, color=EDGE, width=1.25, arrow=False, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    no_shadow(c)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow:
        set_arrow(c)
    return c

def new_slide(title=None, notes=None, number=True):
    s = prs.slides.add_slide(BLANK)
    SLIDE_NO[0] += 1
    if title:
        add_text(s, title, 0.55, 0.32, 12.2, 0.8, size=30, bold=True, color=NAVY,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_line(s, 0.62, 1.12, 3.0, 1.12, color=ORANGE, width=2.5)
    if number:
        add_text(s, str(SLIDE_NO[0]), 12.3, 7.0, 0.8, 0.35, size=11, color=GRAY,
                 align=PP_ALIGN.RIGHT)
        add_text(s, 'Parallelizing Congruence Closure  ·  FMCAD 2026', 0.55, 7.0,
                 6, 0.35, size=11, color=GRAY)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s

# ---------------------------------------------------------------- e-graph diagram
LEAF, GATE, ITE = 4.05, 2.3, 0.55
LW, LH = 0.60, 0.46      # leaf node: name only
GW, GH = 0.88, 0.62      # gate node: operator (large) + label
NODES = {  # id: (label, operator, cx, cy) in diagram-local inches
    'r':  ('r',  None, 0.66, LEAF), 'rp': ('r\u2032', None, 1.34, LEAF),
    's':  ('s',  None, 2.30, LEAF), 'sp': ('s\u2032', None, 2.98, LEAF),
    'c':  ('c',  None, 3.94, LEAF), 'cp': ('c\u2032', None, 4.62, LEAF),
    'u':  ('u',  None, 5.58, LEAF), 'up': ('u\u2032', None, 6.26, LEAF),
    'v':  ('v',  None, 7.22, LEAF), 'vp': ('v\u2032', None, 7.90, LEAF),
    'a1': ('a\u2081', 'AND', 1.25, GATE), 'a2': ('a\u2082', 'AND', 2.45, GATE),
    'x1': ('x\u2081', 'XOR', 6.10, GATE), 'x2': ('x\u2082', 'XOR', 7.30, GATE),
    'm1': ('m\u2081', 'ITE', 3.70, ITE),  'm2': ('m\u2082', 'ITE', 5.15, ITE),
}
EDGES = {'a1': ['r', 's'], 'a2': ['rp', 'sp'], 'x1': ['u', 'v'], 'x2': ['up', 'vp'],
         'm1': ['c', 'a1', 'x1'], 'm2': ['cp', 'a2', 'x2']}
OX, OY = 0.45, 1.55      # default diagram origin

def nsize(nid):
    return (GW, GH) if NODES[nid][1] else (LW, LH)

LAB_UP, LAB_RT = 0.30, 0.37   # room the outside label needs

def lpos(nid):
    """Where the term name sits relative to its node."""
    sym = NODES[nid][1]
    if sym is None:
        return None            # leaves carry their name inside
    return 'right' if sym == 'ITE' else 'above'

LEAF_CLASSES = [(['r', 'rp'], 'leaf'), (['s', 'sp'], 'leaf'), (['c', 'cp'], 'leaf'),
                (['u', 'up'], 'leaf'), (['v', 'vp'], 'leaf')]
A_CLASS = (['a1', 'a2'], 'a')
X_CLASS = (['x1', 'x2'], 'x')
M_CLASS = (['m1', 'm2'], 'm')

def _bbox(ids, pad):
    x0 = min(NODES[i][2] - nsize(i)[0] / 2 for i in ids)
    x1 = max(NODES[i][2] + nsize(i)[0] / 2 + (LAB_RT if lpos(i) == 'right' else 0)
             for i in ids)
    y0 = min(NODES[i][3] - nsize(i)[1] / 2 - (LAB_UP if lpos(i) == 'above' else 0)
             for i in ids)
    y1 = max(NODES[i][3] + nsize(i)[1] / 2 for i in ids)
    return (x0 - pad, y0 - pad, x1 - x0 + 2 * pad, y1 - y0 + 2 * pad)

def draw_egraph(slide, ox=OX, oy=OY, classes=(), highlight=(), groups=(),
                dirty=(), labels=None, dim=(), edge_hl=()):
    """classes: list of (member ids, style key); highlight: node ids with amber
    outline; groups: list of id-lists drawn as dotted purple outlines; dirty:
    indices into classes that get a 'dirty' tag; labels: {class index: text};
    edge_hl: (parent, child) pairs drawn as orange arrows."""
    labels = labels or {}
    edge_hl = set(edge_hl)
    # class boxes
    for ci, (members, key) in enumerate(classes):
        fill, line = CLASS_STYLE[key]
        bx, by, bw, bh = _bbox(members, 0.13)
        add_rect(slide, ox + bx, oy + by, bw, bh, fill=fill, line=line, width=1.5,
                 dash=MSO_LINE.DASH, radius=0.25)
        if ci in labels:
            add_text(slide, labels[ci], ox + bx - 0.05, oy + by - 0.36, bw + 0.4, 0.32,
                     size=12, bold=True, color=line)
        if ci in dirty:
            ty = oy + by + bh - 0.12 if key == 'leaf' else oy + by - 0.14
            add_rect(slide, ox + bx + bw - 0.52, ty, 0.56, 0.26, fill=WHITE,
                     line=RED, width=1.25, radius=0.5, text='dirty', size=9.5, bold=True,
                     color=RED)
    # candidate groups
    for members in groups:
        bx, by, bw, bh = _bbox(members, 0.06)
        add_rect(slide, ox + bx, oy + by, bw, bh, fill=None, line=PURPLE, width=1.75,
                 dash=MSO_LINE.ROUND_DOT, radius=0.3)
    # edges: plain first, highlighted on top
    for wanted in (False, True):
        for parent, kids in EDGES.items():
            px, py = NODES[parent][2], NODES[parent][3]
            ph = nsize(parent)[1]
            for k in kids:
                if ((parent, k) in edge_hl) != wanted:
                    continue
                kx, ky = NODES[k][2], NODES[k][3]
                kh = nsize(k)[1]
                # stop above an outside label so the line never crosses the name
                ktop = ky - kh / 2 - (LAB_UP + 0.01 if lpos(k) == 'above' else 0)
                col = ORANGE if wanted else (LIGHT if (parent in dim or k in dim) else EDGE)
                add_line(slide, ox + px, oy + py + ph / 2, ox + kx, oy + ktop,
                         color=col, width=2.5 if wanted else 1.25, arrow=wanted)
    # nodes
    for nid, (lab, sym, cx, cy) in NODES.items():
        w, h = nsize(nid)
        hl = nid in highlight
        shp = add_rect(slide, ox + cx - w / 2, oy + cy - h / 2, w, h,
                       fill=(AMBER_FILL if hl else WHITE),
                       line=(AMBER if hl else RGBColor(0x4A, 0x4F, 0x57)),
                       width=(2.75 if hl else 1.0), radius=0.2)
        tf = shp.text_frame
        tf.margin_top = tf.margin_bottom = Inches(0.0)
        txtcol = LIGHT if nid in dim else INK
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, sym if sym else lab, 16 if sym else 14, color=txtcol, bold=True)
        where = lpos(nid)
        if where == 'above':
            add_text(slide, lab, ox + cx - 0.45, oy + cy - h / 2 - LAB_UP + 0.01, 0.9, 0.28,
                     size=13, bold=True, color=txtcol, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
        elif where == 'right':
            add_text(slide, lab, ox + cx + w / 2 + 0.02, oy + cy - 0.15, 0.35, 0.3,
                     size=13, bold=True, color=txtcol, anchor=MSO_ANCHOR.MIDDLE)

def side_panel(slide, heading, items, x=8.95, y=1.5, w=4.05, size=16):
    add_text(slide, heading, x, y, w, 0.45, size=18, bold=True, color=NAVY)
    add_bullets(slide, items, x, y + 0.45, w, 5.2, size=size, gap=5)

def legend(slide, entries, x=0.55, y=6.45):
    """entries: list of (kind, text) with kind in {'hl','group','class'}"""
    cx = x
    for kind, text in entries:
        if kind == 'hl':
            add_rect(slide, cx, y + 0.05, 0.36, 0.24, fill=AMBER_FILL, line=AMBER, width=2)
        elif kind == 'group':
            add_rect(slide, cx, y + 0.05, 0.36, 0.24, fill=None, line=PURPLE, width=1.75,
                     dash=MSO_LINE.ROUND_DOT)
        else:
            f, l = CLASS_STYLE['a']
            add_rect(slide, cx, y + 0.05, 0.36, 0.24, fill=f, line=l, width=1.5, dash=MSO_LINE.DASH)
        add_text(slide, text, cx + 0.42, y, 2.4, 0.35, size=12, color=GRAY)
        cx += 0.42 + 0.12 * len(text) + 0.35

# ================================================================ slides
# 1 ---- title
s = new_slide(number=False, notes=(
    'Title. Joint work with Amar Shah, equal contribution. I will cover the problem, '
    'the two algorithms, and implementation; Amar covers the evaluation.'))
add_rect(s, 0, 0, 13.333, 7.5, fill=NAVY, line=None, rounded=False)
add_line(s, 0.9, 3.55, 4.2, 3.55, color=ORANGE, width=3)
add_text(s, 'Parallelizing Congruence Closure', 0.85, 2.2, 11.5, 1.3, size=44, bold=True,
         color=WHITE, anchor=MSO_ANCHOR.BOTTOM)
add_text(s, 'Zachary Kent*   ·   Amar Shah*', 0.85, 3.75, 11, 0.6, size=24, color=WHITE)
add_text(s, 'Carnegie Mellon University', 0.85, 4.35, 11, 0.5, size=20,
         color=RGBColor(0xC9, 0xD3, 0xE6))
add_text(s, 'FMCAD 2026', 0.85, 5.0, 11, 0.5, size=20, color=RGBColor(0xC9, 0xD3, 0xE6))
add_text(s, '* equal contribution', 0.85, 6.6, 6, 0.4, size=13, color=RGBColor(0x9A, 0xA8, 0xC4))
add_text(s, 'github.com/amarshah10/ParallelEgraph', 7.3, 6.6, 5.6, 0.4, size=13,
         color=RGBColor(0x9A, 0xA8, 0xC4), align=PP_ALIGN.RIGHT)

# 2 ---- what is congruence closure
s = new_slide('Congruence closure', notes=(
    'Define the problem: given ground equalities over uninterpreted function terms, '
    'compute the smallest equivalence relation containing them that is closed under '
    'congruence. The tiny example shows transitivity followed by congruence refuting a '
    'disequality. Standard implementation: union-find plus a signature table, i.e. an e-graph.'))
add_bullets(s, [
    '**Input:** terms over uninterpreted functions, equalities E',
    '**Output:** smallest equivalence relation ⊇ E closed under congruence',
    '**Deciding a query** s ≠ t: unsatisfiable iff s ≡ t',
    '**Standard engine:** union-find + signature table (an e-graph)',
    (1, 'Nelson & Oppen 1980, Downey, Sethi & Tarjan 1980, Nieuwenhuis & Oliveras 2007'),
], 0.6, 1.5, 7.0, 4.5, size=20)
# congruence rule
add_rect(s, 0.9, 5.15, 6.2, 1.35, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Congruence rule', 1.05, 5.2, 3, 0.35, size=13, bold=True, color=GRAY)
add_text(s, 's₁ ≡ t₁    …    sₖ ≡ tₖ', 1.05, 5.5, 5.9, 0.4, size=20, align=PP_ALIGN.CENTER)
add_line(s, 1.7, 5.95, 6.4, 5.95, color=INK, width=1.5)
add_text(s, 'f(s₁, …, sₖ) ≡ f(t₁, …, tₖ)', 1.05, 6.0, 5.9, 0.4, size=20, align=PP_ALIGN.CENTER)
# example box
add_rect(s, 7.9, 1.6, 4.9, 3.5, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Example', 8.1, 1.7, 3, 0.4, size=16, bold=True, color=NAVY)
add_text(s, 'a = b        b = c        f(a) ≠ f(c)', 8.1, 2.2, 4.5, 0.5, size=22, bold=True)
add_bullets(s, [
    'transitivity:   a ≡ c',
    'congruence:   f(a) ≡ f(c)',
    'contradicts f(a) ≠ f(c)',
    '**⇒ unsatisfiable**',
], 8.1, 2.9, 4.5, 3.0, size=19, gap=10)

# 3 ---- why it matters
s = new_slide('Where congruence closure runs', notes=(
    'It is a core, performance-critical component in many engines and hence in the '
    'applications built on them. Decades of engineering have gone into sequential '
    'implementations.'))
add_text(s, 'Engines', 0.7, 1.5, 5, 0.45, size=22, bold=True, color=NAVY)
add_bullets(s, [
    '**SMT solvers**: Z3, cvc5 (theory of uninterpreted functions)',
    '**SAT solvers**: Kissat (clausal congruence closure)',
    '**Equality saturation**: egg, egglog',
    '**Type unification**',
], 0.6, 2.0, 5.9, 4, size=19)
add_text(s, 'Applications', 6.95, 1.5, 5, 0.45, size=22, bold=True, color=NAVY)
add_bullets(s, [
    '**Circuit equivalence checking**',
    '**Software verification**: Dafny, Verus',
    '**Hardware design**',
    '**Theorem proving**: Lean grind',
], 6.85, 2.0, 5.9, 4, size=19)
add_rect(s, 0.6, 4.7, 12.1, 1.0, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None,
         text='Performance critical: called millions of times inside a solver run', size=20,
         bold=True, color=NAVY)

# 4 ---- the gap
s = new_slide('Sequential in practice, hard in theory', notes=(
    'Every implementation we know of is sequential. Machines are wide. Theory is '
    'discouraging: congruence closure is P-complete, so a polylog-span algorithm is '
    'unlikely. The question we ask is whether practical instances still parallelize well.'))
add_bullets(s, [
    '**All existing congruence closure implementations are sequential**',
    (1, 'worklist algorithms in the style of Nelson & Oppen'),
    '**Multicore is the default**: our test machine has 96 cores, 192 hardware threads',
    '**P-complete** (Kanellakis & Revesz 1989)',
    (1, 'polylogarithmic span with polynomial work is not expected'),
    (1, 'worst-case instances force a long chain of dependent merges'),
], 0.6, 1.5, 12.0, 4.0, size=21, gap=10)
add_rect(s, 0.6, 5.0, 12.1, 1.2, fill=NAVY, line=None,
         text='Question: do practical instances expose enough independent merges for real speedups?',
         size=21, bold=True, color=WHITE)

# 5 ---- contributions
s = new_slide('Contributions', notes=(
    'Three contributions plus open-source code. Emphasize that the algorithms replace the '
    'sequential worklist with bulk-synchronous rounds.'))
rows = [
    ('1', 'Two parallel algorithms: ParentCC and FilterCC',
     'bulk-synchronous rounds of data-parallel primitives over a lock-free concurrent union-find'),
    ('2', 'Workload characterization: congruence depth and width',
     'width, the number of independent merges per round, is the dominant factor empirically'),
    ('3', 'Evaluation on random, synthetic, and circuit-equivalence benchmarks',
     'near-linear speedup up to 32 cores over a tuned sequential baseline'),
    ('4', 'Open source implementation and benchmarks',
     'github.com/amarshah10/ParallelEgraph'),
]
y = 1.55
for num, head, sub in rows:
    add_rect(s, 0.7, y + 0.05, 0.62, 0.62, fill=ORANGE, line=None, radius=0.5, text=num,
             size=20, bold=True, color=WHITE)
    add_text(s, head, 1.55, y - 0.02, 11, 0.5, size=22, bold=True, color=NAVY)
    add_text(s, sub, 1.55, y + 0.45, 11, 0.5, size=17, color=GRAY)
    y += 1.3

# 6 ---- setting: miters and the term grammar
s = new_slide('Setting: combinational equivalence checking', notes=(
    'Where the running example comes from. A miter feeds the same inputs to two circuits, '
    'XORs the corresponding outputs, and asserts the result is 1. The circuits agree on '
    'every input exactly when that assertion is unsatisfiable. Encoding a miter to CNF '
    'destroys the gate structure, so Biere et al. recover the gates and close them under '
    'congruence, which often decides the instance outright. Gates become terms under this '
    'grammar: input wires are leaves, every gate is a function symbol applied to its inputs.'))
# --- miter block diagram ---
add_rect(s, 2.05, 5.72, 3.1, 0.45, fill=RGBColor(0xEE, 0xF0, 0xF3), line=RGBColor(0x9A, 0xA0, 0xA8),
         width=1.25, text='shared inputs  r, s, u, v, c', size=13)
CB = RGBColor(0xDC, 0xEA, 0xF8)
add_rect(s, 1.55, 4.35, 1.7, 0.85, fill=CB, line=RGBColor(0x2E, 0x6F, 0xB0), width=1.5,
         text='**C\u2081**', size=19, color=NAVY)
add_rect(s, 3.95, 4.35, 1.7, 0.85, fill=CB, line=RGBColor(0x2E, 0x6F, 0xB0), width=1.5,
         text='**C\u2082**', size=19, color=NAVY)
add_line(s, 2.40, 5.72, 2.40, 5.20, arrow=True)
add_line(s, 4.80, 5.72, 4.80, 5.20, arrow=True)
add_rect(s, 2.80, 2.95, 1.6, 0.72, fill=RGBColor(0xFB, 0xE4, 0xD3), line=ORANGE, width=1.5,
         text='**XOR**', size=17, color=NAVY)
add_line(s, 2.40, 4.35, 3.20, 3.67, arrow=True)
add_line(s, 4.80, 4.35, 4.00, 3.67, arrow=True)
add_text(s, 'm\u2081', 2.38, 3.85, 0.5, 0.3, size=14, bold=True, color=GRAY)
add_text(s, 'm\u2082', 4.52, 3.85, 0.5, 0.3, size=14, bold=True, color=GRAY)
add_line(s, 3.60, 2.95, 3.60, 2.45, arrow=True)
add_text(s, 'p', 3.30, 2.08, 0.5, 0.35, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_rect(s, 4.00, 2.05, 1.85, 0.42, fill=NAVY, line=None, text='assert p = 1', size=14,
         bold=True, color=WHITE)
# --- right column ---
add_bullets(s, [
    '#Miter',
    'two circuits, same inputs',
    'XOR the corresponding outputs',
    'p = 1  \u27fa  the circuits disagree',
    '**equivalent  \u27fa  p = 1 unsatisfiable**',
], 6.25, 1.42, 6.5, 2.2, size=18, gap=5)
add_rect(s, 6.35, 3.45, 6.35, 1.45, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Terms', 6.55, 3.52, 3, 0.35, size=15, bold=True, color=NAVY)
gb = s.shapes.add_textbox(Inches(6.55), Inches(3.86), Inches(6.0), Inches(1.0))
gtf = gb.text_frame
gtf.word_wrap = False
for i, (code_, note) in enumerate([
        ('t  ::=  x', 'leaf: an input wire'),
        ('   |   f(t\u2081, \u2026, t\u2096)', 'gate: f \u2208 F,  arity k'),
        ('F = {AND, OR, XOR, ITE, \u2026}', '')]):
    par = gtf.paragraphs[0] if i == 0 else gtf.add_paragraph()
    par.space_after = Pt(2)
    add_runs(par, code_.ljust(22), 14, font=MONO)
    if note:
        add_runs(par, note, 13, color=GRAY)
add_bullets(s, [
    '#Why congruence closure',
    'real miters: many distinct gates, same function',
    'CNF encoding discards that structure',
    'Biere et al. 2024: recover gates, close under congruence',
], 6.25, 5.0, 6.5, 1.6, size=17, gap=5)

# 7 ---- running example intro
s = new_slide('Running example: circuit equivalence checking', notes=(
    'Miter adapted from Biere et al., Clausal Congruence Closure, SAT 2024. Two copies of '
    'a three-gate circuit. Each copy has its own input wires; the equalities tie them '
    'together. Assert the outputs differ. Unsat means the circuits are equivalent. In the '
    'diagram, nodes are terms (gates), edges point to children.'))
draw_egraph(s)
side_panel(s, 'Miter of two copies of a circuit', [
    'a₁ = AND(r, s)    x₁ = XOR(u, v)',
    'm₁ = ITE(c, a₁, x₁)',
    'copy 2 identical over r′, s′, u′, v′, c′',
    '#Equalities E',
    'r = r′,  s = s′,  u = u′,  v = v′,  c = c′',
    '#Query',
    'm₁ ≠ m₂  (outputs differ)',
    'unsatisfiable  ⟺  circuits equivalent',
], size=16)
add_text(s, 'nodes: terms (gates)   ·   edges: parent → child', 0.5, 6.45, 8.3, 0.35,
         size=12, color=GRAY)

# 8 ---- reading the diagram: one node, one term
FOCUS = ['m2', 'cp', 'a2', 'x2']
s = new_slide('Reading the diagram: one node, one term', notes=(
    'Focus on a single node before the trace starts. The box carries the operator, the name '
    'sits alongside, and the outgoing edges are its children in argument order. Everything '
    'else on the slide is greyed out so only this term and its three children are in view.'))
draw_egraph(s, highlight=['m2'], edge_hl=[('m2', 'cp'), ('m2', 'a2'), ('m2', 'x2')],
            dim=[n for n in NODES if n not in FOCUS])
add_rect(s, 6.55, 1.28, 2.3, 0.74, fill=AMBER_FILL, line=AMBER, width=1.5)
add_text(s, 'm\u2082 = ITE(c\u2032, a\u2082, x\u2082)', 6.58, 1.33, 2.24, 0.32, size=14.5,
         bold=True, align=PP_ALIGN.CENTER)
add_text(s, 'orange arrows: its 3 children', 6.58, 1.65, 2.24, 0.3, size=11.5, color=GRAY,
         align=PP_ALIGN.CENTER)
add_line(s, 6.55, 1.62, OX + 5.15 + 0.20, OY + ITE - GH / 2 - 0.02, color=AMBER, width=2.25,
         arrow=True)
side_panel(s, 'Notation', [
    '**operator** inside the node',
    '**term name** alongside it',
    '**one edge per child**, in argument order',
    'leaves: circuit inputs, no children',
    '#Same term, three ways',
    'gate m\u2082 in the circuit',
    'term ITE(c\u2032, a\u2082, x\u2082)',
    'node with 3 outgoing edges',
])

# 9 ---- what the hand trace is doing
s = new_slide('What we do by hand', notes=(
    'Before the trace, state the procedure. Put every input equality into the union-find, '
    'then repeatedly look for two terms with the same operator whose children are pairwise '
    'in the same class, merge them, and re-examine the parents of whatever just merged. '
    'Stop when no pair matches. That fixpoint is the congruence closure. The next four '
    'slides run exactly this loop on the miter, one merge per slide.'))
steps = [
    ('1', 'Union every input equality',
     'r \u2261 r\u2032,  s \u2261 s\u2032,  u \u2261 u\u2032,  v \u2261 v\u2032,  c \u2261 c\u2032'),
    ('2', 'Look for a congruent pair',
     'same operator, children pairwise equivalent'),
    ('3', 'Merge their two classes',
     'one merge at a time, in whatever order we pick'),
    ('4', 'Re-examine the parents of what just merged',
     'a merge can make new pairs congruent'),
    ('5', 'Stop when no pair matches',
     'that fixpoint is the congruence closure'),
]
y = 1.55
for num, head, sub in steps:
    add_rect(s, 0.7, y + 0.02, 0.55, 0.55, fill=ORANGE, line=None, radius=0.5, text=num,
             size=18, bold=True, color=WHITE)
    add_text(s, head, 1.45, y - 0.04, 6.6, 0.5, size=19, bold=True, color=NAVY)
    add_text(s, sub, 1.45, y + 0.38, 6.6, 0.4, size=15, color=GRAY)
    y += 0.97
add_rect(s, 8.35, 1.55, 4.35, 3.15, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Test in step 2', 8.55, 1.65, 4.0, 0.4, size=17, bold=True, color=NAVY)
add_text(s, 'f(M\u2081, \u2026, M\u2099)   and   f(N\u2081, \u2026, N\u2099)', 8.5, 2.1, 4.05, 0.4,
         size=16, align=PP_ALIGN.CENTER)
add_text(s, 'congruent when', 8.5, 2.52, 4.05, 0.32, size=13, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, 'Find(M\u1d62) = Find(N\u1d62)  for every i', 8.5, 2.86, 4.05, 0.4, size=16,
         align=PP_ALIGN.CENTER)
add_bullets(s, [
    'classes live in a union-find',
    'Find(t) = representative of t',
], 8.45, 3.4, 4.15, 1.2, size=15, gap=6)
add_rect(s, 8.35, 4.95, 4.35, 1.35, fill=NAVY, line=None,
         text='Next: this loop on the miter,\none merge per slide', size=17, bold=True, color=WHITE)

# 10-13 ---- sequential walkthrough
s = new_slide('Congruence closure by hand: step 0', notes=(
    'First record the input equalities. The dashed boxes are equivalence classes.'))
draw_egraph(s, classes=LEAF_CLASSES)
side_panel(s, 'Step 0: input equalities', [
    'Union(r, r′), Union(s, s′), Union(u, u′), Union(v, v′), Union(c, c′)',
    'dashed box = equivalence class',
    'every gate still in its own class',
])
legend(s, [('class', 'equivalence class')])

s = new_slide('Congruence closure by hand: step 1', notes=(
    'a1 and a2 have the same symbol and pairwise-equivalent children, so they merge.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS], highlight=['a1', 'a2'])
side_panel(s, 'Step 1: the AND gates', [
    'a₁ = AND(r, s)     a₂ = AND(r′, s′)',
    'same symbol, children pairwise equivalent',
    '**congruence ⇒ a₁ ≡ a₂**',
    'Union(a₁, a₂)',
])
legend(s, [('hl', 'terms compared'), ('class', 'equivalence class')])

s = new_slide('Congruence closure by hand: step 2', notes=('Same for the XOR gates.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS], highlight=['x1', 'x2'])
side_panel(s, 'Step 2: the XOR gates', [
    'x₁ = XOR(u, v)     x₂ = XOR(u′, v′)',
    '**congruence ⇒ x₁ ≡ x₂**',
    'Union(x₁, x₂)',
])
legend(s, [('hl', 'terms compared'), ('class', 'equivalence class')])

s = new_slide('Congruence closure by hand: step 3', notes=(
    'Now the ITE gates have pairwise-equivalent children, so they merge, refuting the '
    'query. Note that step 3 depended on steps 1 and 2, but 1 and 2 did not depend on each other.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS, M_CLASS], highlight=['m1', 'm2'])
side_panel(s, 'Step 3: the ITE gates', [
    'm₁ = ITE(c, a₁, x₁)     m₂ = ITE(c′, a₂, x₂)',
    'children now pairwise equivalent',
    '**congruence ⇒ m₁ ≡ m₂**',
    'contradicts query m₁ ≠ m₂',
    '**⇒ circuits equivalent**',
])
legend(s, [('hl', 'terms compared'), ('class', 'equivalence class')])

# 14 ---- sequential baseline
s = new_slide('Sequential baseline: the worklist algorithm', notes=(
    'What we just did by hand is the classical worklist algorithm, in the variant of '
    'Nieuwenhuis and Oliveras. The signature table detects congruences in O(1). Each '
    'merge pushes the loser\'s parents back onto the worklist. We compared three sequential '
    'variants and this was the fastest, so it is our baseline.'))
add_bullets(s, [
    '**Union-find** over all terms',
    '**Signature table**: hash of (symbol, Find(child₁), …, Find(childₖ)) → representative',
    '**Parents[class]**: terms with a child in the class',
    '**FIFO worklist** of compound terms, seeded in reverse-topological order',
    '**Loop**: pop e; on a signature collision with e′, Union(e, e′); push Parents[loser]',
    'one merge at a time; each merge re-examines the parents of the losing class',
], 0.6, 1.5, 7.6, 5, size=19, gap=8)
add_rect(s, 8.6, 1.6, 4.2, 3.4, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Sequential variants we tried', 8.8, 1.7, 3.9, 0.4, size=15, bold=True, color=NAVY)
add_bullets(s, [
    '**Worklist** (Nieuwenhuis & Oliveras): fastest, used as baseline',
    '**Topological-sort fixpoint**: repeated full passes',
    '**Downey–Sethi–Tarjan**: O(n log n) work, hashtable instead of trie',
], 8.65, 2.15, 4.1, 4, size=15, gap=8)

# 15 ---- observation: depth and width
s = new_slide('Observation: independent merges form rounds', notes=(
    'Steps 1 and 2 did not depend on each other; they could be done concurrently, as one '
    'round. Step 3 forms a second round. Depth is the number of rounds; width is the max '
    'number of merges in a round. Here both are 2. On real benchmarks width reaches millions, '
    'and that is the parallelism we exploit.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS, M_CLASS],
            labels={5: 'round 1', 6: 'round 1', 7: 'round 2'})
add_text(s, 'round 0: input equalities', 0.45, OY + LEAF + LH / 2 + 0.32, 8.55, 0.35, size=12,
         bold=True, color=GRAY, align=PP_ALIGN.CENTER)
side_panel(s, 'Depth and width', [
    'steps 1 and 2 are independent: one round',
    'step 3 depends on both: next round',
    '#Congruence depth',
    'number of rounds until fixpoint  (here 2)',
    '#Congruence width',
    'max new merges in one round  (here 2)',
    '#In practice',
    'width reaches thousands to millions',
    'each round: many independent unions',
])

# 16 ---- building blocks
s = new_slide('Two building blocks from parallel algorithms', notes=(
    'Concurrent union-find from Alistarh et al.: simple CAS-based variant, since fancier '
    'ones did not perform better in their study. Semisort from Gu et al. groups by key '
    'without a total order, linear expected work. Both come via ParlayLib.'))
add_rect(s, 0.6, 1.5, 6.0, 4.6, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Concurrent union-find', 0.8, 1.6, 5.6, 0.45, size=21, bold=True, color=NAVY)
add_text(s, 'Alistarh, Fedorov & Koval, OPODIS 2019', 0.8, 2.02, 5.6, 0.35, size=13, color=GRAY)
add_bullets(s, [
    'array UF[i]: parent index, or rank if root',
    '**Find**: follow parents, path-compress with CAS',
    '**Union**: link lower rank under higher via CAS; retry on failure',
    'best-effort linking by rank',
    'lock-free and linearizable',
    'contention = failed CAS + retry, never blocking',
], 0.7, 2.4, 5.8, 4, size=17, gap=7)
add_rect(s, 6.85, 1.5, 5.9, 4.6, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Semisort (GroupBy)', 7.05, 1.6, 5.5, 0.45, size=21, bold=True, color=NAVY)
add_text(s, 'Gu, Shun, Sun & Blelloch, SPAA 2015', 7.05, 2.02, 5.5, 0.35, size=13, color=GRAY)
add_bullets(s, [
    'input: (key, value) pairs',
    'output: equal keys adjacent, no total order across keys',
    '**O(n) work, O(log n) span** in expectation',
    'key = signature hash: (symbol, Find(child₁), …, Find(childₖ))',
    'one call gathers all candidate congruences of a round',
], 6.95, 2.4, 5.7, 4, size=17, gap=7)
add_text(s, 'Glue: ParlayLib (Blelloch, Anderson & Dhulipala 2020) for parfor, scheduler, '
            'integer sort, filter.  Rounds are bulk-synchronous (Valiant 1990).',
         0.6, 6.5, 12.2, 0.5, size=14, color=GRAY)

# 17 ---- ParentCC overview
s = new_slide('ParentCC: a bulk-synchronous closure loop', notes=(
    'Replace the worklist with rounds. Work holds the pre-merge roots of the classes merged in the '
    'previous round. Fold their parent lists into the new roots, take the parents as the '
    'frontier, group the frontier by signature with semisort, merge every group that spans '
    'more than one class, and all pre-merge roots of merged groups form the next Work. Stop when Work is empty.'))
add_rect(s, 0.6, 1.45, 12.1, 0.7, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None,
         text='Round 0:  parfor (u = v) ∈ E: Union(u, v)        seed Frontier ← all compound terms', size=18,
         color=INK)
BW, BH = 3.7, 1.35
bx1, bx2 = 1.2, 8.4
by1, by2 = 2.55, 4.95
b1 = add_rect(s, bx1, by1, BW, BH, fill=RGBColor(0xDC, 0xEA, 0xF8), line=RGBColor(0x2E, 0x6F, 0xB0),
              width=1.5, text='**Work**\npre-merge roots of the classes\nmerged in the previous round', size=16)
b2 = add_rect(s, bx2, by1, BW, BH, fill=RGBColor(0xDC, 0xEA, 0xF8), line=RGBColor(0x2E, 0x6F, 0xB0),
              width=1.5, text='**Frontier**\nseed round: all compound terms\nlater: parents of the changed classes', size=16)
b3 = add_rect(s, bx2, by2, BW, BH, fill=RGBColor(0xEA, 0xE2, 0xF5), line=PURPLE,
              width=1.5, text='**GroupBy signature**\nsemisort the frontier by\n(symbol, Find(children))', size=16)
b4 = add_rect(s, bx1, by2, BW, BH, fill=RGBColor(0xFB, 0xE4, 0xD3), line=ORANGE,
              width=1.5, text='**Merge**\nparallel unions within each group\nspanning more than one class', size=16)
add_line(s, bx1 + BW, by1 + BH / 2, bx2, by1 + BH / 2, color=INK, width=1.75, arrow=True)
add_line(s, bx2 + BW / 2, by1 + BH, bx2 + BW / 2, by2, color=INK, width=1.75, arrow=True)
add_line(s, bx2, by2 + BH / 2, bx1 + BW, by2 + BH / 2, color=INK, width=1.75, arrow=True)
add_line(s, bx1 + BW / 2, by2, bx1 + BW / 2, by1 + BH, color=INK, width=1.75, arrow=True)
add_text(s, 'barrier', 6.3, by1 + 0.22, 1.5, 0.35, size=12, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, 'barrier', 6.3, by2 + 0.22, 1.5, 0.35, size=12, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, 'pre-merge roots of\nmerged groups → next Work', bx1 + BW / 2 + 0.12, 4.0, 2.4, 0.7, size=12, color=GRAY)
add_text(s, 'repeat until Work = ∅\nresult: union-find holds the congruence classes',
         5.25, 3.95, 3.2, 1.0, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# 18-23 ---- ParentCC on the example
FRONT1 = ['a1', 'a2', 'x1', 'x2', 'm1', 'm2']
s = new_slide('ParentCC on the example: round 0', notes=(
    'All five input unions run in parallel. Work is seeded with the terms of the equalities.'))
draw_egraph(s, classes=LEAF_CLASSES)
side_panel(s, 'Round 0', [
    '**parfor** (u = v) ∈ E: Union(u, v)',
    'five unions, all concurrent',
    'Work ← {r, r′, s, s′, u, u′, v, v′, c, c′}',
])
legend(s, [('class', 'class')])

s = new_slide('ParentCC on the example: round 1, frontier', notes=(
    'Fold the parent lists of the dethroned leaves into their new roots. The frontier is '
    'every parent of a changed class, which includes the ITE gates because c changed. '
    'Semisort groups by signature: the ANDs match, the XORs match, but the ITEs do not yet '
    'because a1 and a2 are still in different classes.'))
draw_egraph(s, classes=LEAF_CLASSES, highlight=FRONT1,
            groups=[['a1', 'a2'], ['x1', 'x2'], ['m1'], ['m2']])
side_panel(s, 'Round 1: frontier and groups', [
    'fold: Parents[r] ∪= Parents[r′] = {a₁, a₂}, …',
    'Frontier = {a₁, a₂, x₁, x₂, m₁, m₂}',
    (1, 'seed round: every compound term'),
    'GroupBy(Congruent, Frontier):',
    (1, '{a₁, a₂}   AND over classes (r, s)'),
    (1, '{x₁, x₂}   XOR over classes (u, v)'),
    (1, '{m₁}  {m₂}   a₁, a₂ not yet equivalent'),
])
legend(s, [('hl', 'frontier'), ('group', 'signature group'), ('class', 'class')])

s = new_slide('ParentCC on the example: round 1, merge', notes=(
    'Both spanning groups merge concurrently. The losing representatives become the next Work.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS], highlight=['a1', 'a2', 'x1', 'x2'])
side_panel(s, 'Round 1: merge', [
    '**parfor** over groups spanning > 1 class:',
    (1, 'Union(a₁, a₂)  ∥  Union(x₁, x₂)'),
    'width of this round: 2',
    'pre-merge roots of both groups: a₁, a₂, x₁, x₂',
    'Work ← {a₁, a₂, x₁, x₂}',
])
legend(s, [('hl', 'merged this round'), ('class', 'class')])

s = new_slide('ParentCC on the example: round 2, frontier', notes=(
    'Fold the parents of a2 and x2 into a1 and x1. The frontier is the two ITEs, and now '
    'their signatures agree.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS], highlight=['m1', 'm2'],
            groups=[['m1', 'm2']])
side_panel(s, 'Round 2: frontier and groups', [
    'fold: Parents[a₁] ∪= Parents[a₂] = {m₁, m₂}',
    'fold: Parents[x₁] ∪= Parents[x₂] = {m₁, m₂}',
    'Frontier = {m₁, m₂}',
    'GroupBy: {m₁, m₂}',
    (1, 'ITE over classes (c, a₁, x₁)'),
])
legend(s, [('hl', 'frontier'), ('group', 'signature group'), ('class', 'class')])

s = new_slide('ParentCC on the example: round 2, merge', notes=('One union; both pre-merge roots enter Work.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS, M_CLASS], highlight=['m1', 'm2'])
side_panel(s, 'Round 2: merge', [
    'Union(m₁, m₂)',
    'width of this round: 1',
    'Work ← {m₁, m₂}',
])
legend(s, [('hl', 'merged this round'), ('class', 'class')])

s = new_slide('ParentCC on the example: round 3, fixpoint', notes=(
    'm2 has no parents, so the frontier is empty and Work becomes empty. Three rounds '
    'instead of eight sequential merges, and the query is refuted.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS, M_CLASS])
side_panel(s, 'Round 3: done', [
    'fold: Parents[m₂] = ∅; m₁ is its own root',
    'Frontier = ∅,  Work = ∅',
    '**fixpoint reached**',
    '#Result',
    'm₁ ≡ m₂ contradicts m₁ ≠ m₂',
    '**circuits equivalent**',
    '#Cost',
    '3 rounds (widths 5, 2, 1)',
    'vs. 8 sequential unions',
])
legend(s, [('class', 'class')])

# 24 ---- pseudocode
s = new_slide('ParentCC pseudocode', notes=(
    'The full loop. Phases are separated by barriers, so signatures are computed against a '
    'union-find that no thread is modifying. MergeCongruenceClass unions a group by '
    'divide and conquer, in parallel.'))
code = [
    ('ParentCC(E):', 0),
    ('parfor (u = v) ∈ E:  Union(u, v)', 1),
    ('Work ← terms of E', 1),
    ('while Work ≠ ∅:', 1),
    ('parfor c ∈ Work with c ≠ Find(c):                 ▹ fold parent lists', 2),
    ('Parents[Find(c)] ← Parents[Find(c)] ∪ Parents[c]', 3),
    ('Frontier ← all compound terms          on round 0   ▹ seed round', 2),
    ('Frontier ← ⋃ Parents[Find(c)] for c ∈ Work   otherwise', 2),
    ('groups ← GroupBy(Congruent, Frontier)          ▹ semisort on sigs', 2),
    ('Work ← ∅', 2),
    ('parfor g ∈ groups spanning > 1 class:', 2),
    ('rep ← MergeCongruenceClass(g)                    ▹ parallel unions', 3),
    ('Work ← Work ∪ pre-merge roots of g', 3),
    ('', 0),
    ('MergeCongruenceClass(S):                             ▹ divide and conquer', 0),
    ('if |S| = 1: return the element', 1),
    ('l, r ← MergeCongruenceClass(S[0..n/2]) ∥ MergeCongruenceClass(S[n/2..n])', 1),
    ('return Union(l, r)', 1),
]
add_rect(s, 0.6, 1.45, 8.9, 5.35, fill=RGBColor(0xF7, 0xF7, 0xF9), line=None, radius=0.05)
tb = s.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(8.6), Inches(5.2))
tf = tb.text_frame
tf.word_wrap = False
for i, (line, ind) in enumerate(code):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(1.5)
    txt = '    ' * ind + line
    if '▹' in txt:
        a, b = txt.split('▹')
        add_runs(p, a, 13.5, font=MONO)
        add_runs(p, '▹' + b, 13.5, font=MONO, color=GRAY)
    else:
        add_runs(p, txt, 13.5, font=MONO, bold=(ind == 0 and line != ''))
add_bullets(s, [
    '**barrier** between fold, group, and merge phases',
    'signatures use a union-find no thread is modifying',
    '**round 0** examines every compound term',
    (1, 'terms can be congruent before any union'),
    '**Congruent(X, Y)**: same symbol, Find(Mᵢ) = Find(Nᵢ) for all i',
    'implemented with ParlayLib parfor, group_by, filter',
], 9.7, 1.5, 3.4, 5, size=15, gap=8)

# 25 ---- correctness
s = new_slide('Why ParentCC is correct', notes=(
    'Four short arguments. Determinism per round follows from linearizability plus '
    'order-independence of the partition generated by a set of unions. Termination: each '
    'productive round reduces the class count. Soundness by the invariant that union-find classes '
    'sit inside congruence classes. Completeness: the last merge among child pairs puts '
    'both parents on the next frontier.'))
blocks = [
    ('Rounds are deterministic',
     ['union-find is linearizable', 'the partition from a set of unions is order independent',
      'barrier: signatures see a fixed union-find']),
    ('Termination',
     ['a spanning group forces at least one successful union',
      'so every productive round reduces the class count',
      'at most |T| − 1 productive rounds']),
    ('Soundness',
     ['invariant: every union-find class ⊆ a class of ≡',
      'grouped terms share symbol and child classes', 'so each merge follows by congruence']),
    ('Completeness',
     ['children already equivalent at round 0: both parents are in the seed frontier',
      'else take the round of the last merge among the children',
      'all its pre-merge roots enter Work, so both parents are in the next frontier',
      'equal signatures there ⇒ merged; classes never shrink']),
]
x = 0.6
for head, items in blocks:
    add_rect(s, x, 1.5, 2.95, 4.0, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
    add_text(s, head, x + 0.12, 1.6, 2.75, 0.5, size=17, bold=True, color=NAVY)
    add_bullets(s, items, x + 0.05, 2.15, 2.85, 4.2, size=14, gap=8)
    x += 3.08

# 26 ---- FilterCC
s = new_slide('FilterCC: drop the parent lists', notes=(
    'Parent lists cost allocations: every fold appends lists. FilterCC keeps one dirty bit '
    'per class instead and recomputes the frontier by filtering all terms for a dirty child. '
    'It scans everything each round, but the scan is a cheap parallel filter with no '
    'allocation. On most workloads that tradeoff wins; Amar will show the numbers.'))
add_bullets(s, [
    '**ParentCC cost**: folding parent lists means frequent allocation and list appends',
    '**FilterCC idea**: one dirty bit per class, set on merge',
    '**round 0**: every compound term dirty, as in ParentCC',
    '**Frontier** = parallel filter over all terms: keep t if some child sits in a dirty class',
    'same frontier as ParentCC, computed without parent lists',
], 0.6, 1.45, 12.1, 2.6, size=19, gap=8)
# comparison table
tx, ty = 0.7, 4.0
cols = [('', 2.3), ('ParentCC', 4.7), ('FilterCC', 4.7)]
rowsT = [
    ('per-class state', 'parent list', 'one dirty bit'),
    ('frontier', 'parents of merged classes, via list folding', 'filter all terms for a dirty child'),
    ('per-round cost', '∝ |frontier| + allocations', '∝ |all terms|, no allocation'),
    ('wins when', 'few terms, many rounds', 'most workloads we measured'),
]
cx = tx
for name, w in cols:
    add_rect(s, cx, ty, w - 0.05, 0.45, fill=NAVY, line=None, rounded=False, text=name, size=15,
             bold=True, color=WHITE)
    cx += w
ry = ty + 0.5
for row in rowsT:
    cx = tx
    for (name, w), val in zip(cols, row):
        add_rect(s, cx, ry, w - 0.05, 0.5, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None,
                 rounded=False, text=val, size=14, bold=(name == ''), color=(NAVY if name == '' else INK))
        cx += w
    ry += 0.55

# 27-28 ---- FilterCC on the example
s = new_slide('FilterCC on the example: round 1', notes=(
    'After the input unions the five leaf classes are dirty. Filter all 16 terms: keep those '
    'with a child in a dirty class. That is the same frontier ParentCC computed, found by a '
    'scan instead of parent lists.'))
draw_egraph(s, classes=LEAF_CLASSES, highlight=FRONT1, dirty=[0, 1, 2, 3, 4],
            groups=[['a1', 'a2'], ['x1', 'x2'], ['m1'], ['m2']])
side_panel(s, 'Round 1', [
    'seed round: **every compound term is dirty**',
    'later rounds filter for a child in a dirty class',
    'Frontier = {a₁, a₂, x₁, x₂, m₁, m₂}',
    'GroupBy, merge spanning groups as before',
    'clear dirty; set dirty[rep] for merged groups',
])
legend(s, [('hl', 'passes filter'), ('group', 'signature group'), ('class', 'class')])

s = new_slide('FilterCC on the example: round 2', notes=(
    'Only the two merged gate classes are dirty. The filter keeps the ITEs. After merging '
    'them, the next filter finds nothing, so the algorithm stops.'))
draw_egraph(s, classes=LEAF_CLASSES + [A_CLASS, X_CLASS], highlight=['m1', 'm2'], dirty=[5, 6],
            groups=[['m1', 'm2']])
side_panel(s, 'Round 2', [
    'dirty = {class of a₁, class of x₁}',
    'filter all terms: Frontier = {m₁, m₂}',
    'GroupBy: {m₁, m₂}  ⇒  Union(m₁, m₂)',
    'dirty = {class of m₁}',
    '#Round 3',
    'filter finds no term with a dirty child',
    'Frontier = ∅  ⇒  **done**',
])
legend(s, [('hl', 'passes filter'), ('group', 'signature group'), ('class', 'class')])

# 29 ---- implementation details
s = new_slide('Implementation details that mattered', notes=(
    'Grouping: ParlayLib integer sort on the hashes followed by a sequential bucket pass '
    'per hash, which beats the library semisort on the small buckets we see. The sequential '
    'baseline got the same care: arity-specialized signature tables and a bump allocator, '
    'so the speedups are measured against a strong baseline.'))
add_rect(s, 0.6, 1.5, 6.0, 4.6, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Parallel algorithms', 0.8, 1.6, 5.6, 0.45, size=21, bold=True, color=NAVY)
add_bullets(s, [
    '**Grouping**: ParlayLib integer sort on signature hashes, then sequential bucketing per hash',
    (1, 'detects hash collisions'),
    (1, 'beats the library semisort on the small buckets typical here'),
    '**Merging a group**: divide-and-conquer parallel unions',
    '**Scheduling**: ParlayLib work-stealing scheduler, parfor, filter',
    '**Union-find**: single array, rank in the high bit, CAS only',
], 0.7, 2.15, 5.8, 4.2, size=16, gap=7)
add_rect(s, 6.85, 1.5, 5.9, 4.6, fill=RGBColor(0xF4, 0xF5, 0xF7), line=None)
add_text(s, 'Sequential baseline', 7.05, 1.6, 5.5, 0.45, size=21, bold=True, color=NAVY)
add_bullets(s, [
    '**Signature tables specialized by arity** 1, 2, 3, 4; general table for the rest',
    (1, 'fixed-size entries, no per-signature allocation'),
    '**Bump allocator** for the general table',
    '**Worklist** seeded in reverse-topological order',
    (1, 'single pass when the quotient DAG is acyclic'),
    'baseline tuned as carefully as the parallel code',
], 6.95, 2.15, 5.7, 4.2, size=16, gap=7)
add_text(s, 'C++, compiled with g++ -O3.  Code and benchmarks: github.com/amarshah10/ParallelEgraph',
         0.6, 6.5, 12.2, 0.4, size=14, color=GRAY)

# 30 ---- handoff divider
s = new_slide(number=False, notes='Hand off to Amar for the evaluation.')
add_rect(s, 0, 0, 13.333, 7.5, fill=NAVY, line=None, rounded=False)
add_line(s, 0.9, 3.55, 4.2, 3.55, color=ORANGE, width=3)
add_text(s, 'Evaluation', 0.85, 2.2, 11.5, 1.3, size=44, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.BOTTOM)
add_text(s, 'random, synthetic, and circuit-equivalence benchmarks', 0.85, 3.75, 11, 0.6,
         size=22, color=RGBColor(0xC9, 0xD3, 0xE6))

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'parallelcc-part1.pptx')
prs.save(out)
print('wrote', out, 'slides:', SLIDE_NO[0])
